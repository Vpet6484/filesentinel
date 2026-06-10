# modules/leak_analyzer.py

import os

from modules.watermark import decode_watermark

from PyPDF2 import PdfReader
from docx import Document
from openpyxl import load_workbook


# =====================================================
# MAIN ANALYZER
# =====================================================

def analyze_file(path):

    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        return analyze_pdf(path)

    elif ext == ".docx":
        return analyze_docx(path)

    elif ext == ".xlsx":
        return analyze_xlsx(path)

    else:
        raise ValueError("Unsupported file format")


# =====================================================
# PDF
# =====================================================

def analyze_pdf(path):

    reader = PdfReader(path)

    meta = reader.metadata

    if not meta or "/FileSentinel-Watermark" not in meta:
        return None

    token = meta["/FileSentinel-Watermark"]

    return decode_watermark(token)


# =====================================================
# DOCX
# =====================================================

def analyze_docx(path):

    doc = Document(path)

    for p in doc.paragraphs:

        text = p.text

        if text.startswith("FS-WM:"):

            token = text.replace("FS-WM:", "").strip()

            return decode_watermark(token)

    return None




# =====================================================
# XLSX
# =====================================================

def analyze_xlsx(path):

    wb = load_workbook(path)

    props = wb.properties

    token = props.keywords

    if not token:
        return None

    return decode_watermark(token)
